"""Makes a small sample .pptx for testing the Weekly Reports conversion. Needs: pip install python-pptx

    python scripts/make-sample-deck.py sample.pptx 18
"""
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "sample.pptx"
    slides = int(sys.argv[2]) if len(sys.argv) > 2 else 6
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9

    for n in range(1, slides + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        slide.shapes.title.text = f"Weekly Performance Review - Slide {n}"
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(5.5), Inches(3))
        tf = box.text_frame
        tf.word_wrap = True
        for line in ("Conversion rate up 12%", "Add-to-cart up 14%", "Checkout starts up 9%"):
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(22)
        data = CategoryChartData()
        data.categories = ["Mon", "Tue", "Wed", "Thu", "Fri"]
        data.add_series("Sessions", (120 + n, 150, 170, 160 + n, 190))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.8), Inches(1.6), Inches(5.8), Inches(4.5), data)

    prs.save(out)
    print(f"wrote {out} ({slides} slides)")


main()
